"""Membuat SATU file template Refinitiv refinitiv/data-refinitiv.xlsx (rumus TR() LSEG Workspace for Excel).

Urutan sheet = urutan kerja: MULAI, config, tickers, bench, data saham 2025-sekarang (P1), statis,
konstituen, tambahan, lalu data saham 2020-2024 dan 2015-2019 (P2).
Daftar ticker ditarik otomatis oleh rumus di sheet tickers (konstituen IHSG terkini).
Mode kalkulasi disimpan Manual supaya ratusan TR() tidak menembak sekaligus saat file dibuka.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.workbook.properties import CalcProperties

import konfigurasi as k

R = k.REF
OUT = k.ROOT / "refinitiv"
BARIS_BENCH = 20  # tabel benchmark di sheet config: A = nama, B = RIC, C = field, D = alternatif
BARIS_TAMBAHAN = BARIS_BENCH + len(R["bench"]) + 3  # tabel tambahan: A = nama, C = field, D = parameter
RIC_IHSG = f"config!$B${BARIS_BENCH}"
SHEET_SAHAM = ["tr_pct", "px_last", "px_volume"]


def blok_tr(sh, b, instrumen, fields, param, lebar, label):
    c = b * lebar + 1
    sh.cell(1, c, label)
    sh.cell(2, c, f"=TR({instrumen},{fields},{param})")


def blok_saham(sh, n_batch, fields, param, lebar):
    """Blok TR berdampingan, masing-masing atas `n_batch` ticker dari sheet tickers kolom B (hasil rumus)."""
    for b in range(-(-R["n_universe"] // n_batch)):
        r0 = 2 + b * n_batch
        r1 = r0 + n_batch - 1
        blok_tr(sh, b, f"tickers!$B${r0}:$B${r1}", fields, param, lebar, f"tickers B{r0}:B{r1}")


def blok_tumpuk(sh, tanggal, ric, fields, jarak):
    """Snapshot ditumpuk ke bawah: label 'snapshot <tgl>' lalu TR() di baris berikutnya."""
    for i, t in enumerate(tanggal):
        r = 1 + i * jarak
        sh.cell(r, 1, f"snapshot {t}")
        param = '"CH=Fd RH=IN"' if t == "terkini" else f'"SDate={t} CH=Fd RH=IN"'
        sh.cell(r + 1, 1, f'=TR({ric},"{fields}",{param})')


def sheet_mulai(wb, urutan):
    sh = wb.active
    sh.title = "MULAI"
    sh.column_dimensions["A"].width = 120
    baris = [
        ("CARA PAKAI FILE INI", True),
        ("", False),
        ("1. TES ADD-IN: ketik di sel D3 sheet ini:  =TR(\".JKSE\",\"TR.PriceClose\")  lalu Enter. Harus keluar angka (harga IHSG).", False),
        ("   Kalau #NAME?, coba =RDP.Data(\".JKSE\",\"TR.PriceClose\"). Kalau itu yang jalan: Ctrl+H, ganti semua '=TR(' menjadi '=RDP.Data(' (Within: Workbook).", False),
        ("", False),
        ("2. CARA REFRESH SATU SHEET: klik tab sheet -> menu Formulas -> Calculate Sheet (atau Fn+Shift+F9).", False),
        ("   JANGAN klik 'Calculate Now' (menarik semua sheet sekaligus, Excel bisa macet). Tunggu sampai sel berhenti loading, lalu Ctrl+S.", False),
        ("", False),
        ("3. (Opsional, 20 menit) Sheet config: isi RIC/field yang kosong; kalau sebuah deret gagal, coba isian kolom 'Alternatif jika kosong'.", False),
        ("   Kalau tidak ketemu, biarkan kosong - aman, nanti dilewati.", False),
        ("", False),
        ("4. REFRESH SHEET SESUAI URUTAN INI (Save setiap selesai satu sheet):", True),
    ]
    baris += [(f"   {i}. {nama}{'   <- P1, wajib untuk brief' if p1 else ''}", p1) for i, (nama, p1) in enumerate(urutan, start=1)]
    baris += [
        ("", False),
        ("   Kalau sheet 'tickers' kosong setelah refresh: tempel daftar kode saham (format BBCA.JK) di kolom B mulai B2.", False),
        ("   Sheet saham 2020 dan 2015 paling berat; boleh dilanjutkan di kunjungan lain.", False),
        ("", False),
        ("5. SELESAI: Save, bawa pulang lewat flashdisk / Google Drive (file besar), taruh di 10-riset-msci\\data\\mentah\\data-refinitiv.xlsx", True),
        ("   Ada error yang aneh? Screenshot dan kirim ke Claude.", False),
    ]
    for i, (teks, tebal) in enumerate(baris, start=1):
        c = sh.cell(i, 1, teks)
        c.font = Font(bold=tebal, size=12 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=False)


def isi_config(cfg):
    for i, (label, nilai) in enumerate([
        ("Mulai deret pasar", R["pasar_mulai"]), ("Akhir", R["akhir"]),
        ("RIC MSCI Indonesia", R["ric_msci_id"]), ("RIC FTSE Indonesia (isi dulu)", R["ric_ftse_id"]),
        ("Mulai data statis historis", R["statis_mulai"]),
    ], start=1):
        cfg.cell(i, 1, label)
        cfg.cell(i, 2, nilai)
    cfg["A7"], cfg["B7"] = "Parameter deret pasar harian", '="SDate="&$B$1&" EDate="&$B$2&" Frq=D CH=Fd RH=IN"'
    cfg["A8"], cfg["B8"] = "Parameter statis kuartalan", '="SDate="&$B$5&" EDate="&$B$2&" Frq=CQ CH=Fd RH=IN"'
    for i, (suf, mulai, akhir) in enumerate(R["periode_saham"]):
        r = 10 + i
        cfg.cell(r, 1, f"Periode saham {suf}")
        cfg.cell(r, 2, f'="SDate={mulai} EDate={akhir} Frq=D CH=Fd RH=IN"')
    for j, h in enumerate(["Deret pasar", "RIC (isi/cek)", "Field (ganti bila kosong)", "Alternatif jika kosong"], start=1):
        cfg.cell(BARIS_BENCH - 1, j, h)
    for i, (nama, ric, field, alt) in enumerate(R["bench"]):
        for j, v in enumerate([nama, ric, field, alt], start=1):
            cfg.cell(BARIS_BENCH + i, j, v)
    for j, h in enumerate(["Data tambahan", "", "Field (isi dulu)", "Parameter"], start=1):
        cfg.cell(BARIS_TAMBAHAN - 1, j, h)
    for i, (nama, field, param) in enumerate(R["tambahan"]):
        cfg.cell(BARIS_TAMBAHAN + i, 1, nama)
        cfg.cell(BARIS_TAMBAHAN + i, 3, field)
        cfg.cell(BARIS_TAMBAHAN + i, 4, param)
    cfg.column_dimensions["A"].width = 30
    cfg.column_dimensions["B"].width = 22
    cfg.column_dimensions["C"].width = 26
    cfg.column_dimensions["D"].width = 70


def sheet_saham(wb, suf):
    baris_param = 10 + [p for p, *_ in R["periode_saham"]].index(suf)
    for s in SHEET_SAHAM:
        field = R["sheet_harga"][s]
        blok_saham(wb.create_sheet(f"{s}_{suf}"), R["batch"], f'"{field}.date;{field}"', f"config!$B${baris_param}",
                   R["lebar_blok"])


def main():
    wb = Workbook()
    p1_suf = R["periode_saham"][0][0]
    lama = [suf for suf, *_ in R["periode_saham"][1:]]
    nama_tambahan = [n for n, _, _ in R["tambahan"]]
    urutan = ([("tickers", True), ("bench", True)] + [(f"{s}_{p1_suf}", True) for s in SHEET_SAHAM]
              + [("static", True), ("msci", True), ("ftse", False), ("universe", False), ("ff_hist", False),
                 ("sh_hist", False)] + [(n, False) for n in nama_tambahan]
              + [(f"{s}_{suf}", False) for suf in lama for s in SHEET_SAHAM])
    sheet_mulai(wb, urutan)
    isi_config(wb.create_sheet("config"))

    tk = wb.create_sheet("tickers")
    tk["A1"] = f'=TR({RIC_IHSG},"TR.IndexConstituentRIC","CH=Fd RH=IN")'

    bn = wb.create_sheet("bench")
    for b, (nama, *_) in enumerate(R["bench"]):
        baris = BARIS_BENCH + b
        blok_tr(bn, b, f"config!$B${baris}", f'config!$C${baris}&".date;"&config!$C${baris}', "config!$B$7",
                R["lebar_blok"], nama)

    sheet_saham(wb, p1_suf)

    blok_saham(wb.create_sheet("static"), R["batch_statis"], f'"{";".join(R["field_statis"])}"', '"CH=Fd RH=IN"',
               len(R["field_statis"]) + 2)
    tanggal = k.tanggal_snapshot()
    indeks = "TR.IndexConstituentRIC;TR.IndexConstituentWeightPercent;TR.IndexConstituentName"  # nama: cadangan bila RIC "NULL" (MSCI)
    blok_tumpuk(wb.create_sheet("msci"), tanggal, "config!$B$3", indeks, R["jarak_snapshot_indeks"])
    blok_tumpuk(wb.create_sheet("ftse"), tanggal, "config!$B$4", indeks, R["jarak_snapshot_indeks"])
    blok_tumpuk(wb.create_sheet("universe"), tanggal + ["terkini"], RIC_IHSG, "TR.IndexConstituentRIC",
                R["jarak_snapshot_ihsg"])
    for nama, field in [("ff_hist", "TR.FreeFloatPct"), ("sh_hist", "TR.SharesOutstanding")]:
        blok_saham(wb.create_sheet(nama), R["batch_statis_hist"], f'"{field}.date;{field}"', "config!$B$8",
                   R["lebar_blok"])
    for i, nama in enumerate(nama_tambahan):
        baris = BARIS_TAMBAHAN + i
        blok_saham(wb.create_sheet(nama), R["batch_statis"],
                   f'config!$C${baris}&".date;"&config!$C${baris}', f"config!$D${baris}", R["lebar_blok"])
    for suf in lama:
        sheet_saham(wb, suf)

    for sh in wb.worksheets[2:]:
        sh.freeze_panes = "A2"
    wb.calculation = CalcProperties(calcMode="manual", fullCalcOnLoad=False, calcOnSave=False)
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / f"{R['nama_file']}.xlsx")
    print(f"Template tersimpan: refinitiv/{R['nama_file']}.xlsx ({len(wb.sheetnames)} sheet)")


if __name__ == "__main__":
    main()
